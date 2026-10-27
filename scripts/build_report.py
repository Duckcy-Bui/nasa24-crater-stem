"""Đức — S3-12: ghép báo cáo từ docs/bao_cao/*.md thành DOCX và PDF.

Định dạng theo yêu cầu môn: A4, Times New Roman 13, giãn dòng 1,3. Lề trái
3 cm, phải/trên/dưới 2 cm, số trang ở chân trang. Mỗi chương bắt đầu trang mới.

Cách chạy:
    pip install -r requirements-report.txt
    python scripts/build_report.py            # DOCX, và PDF nếu tìm thấy LibreOffice
    python scripts/build_report.py --no-pdf   # chỉ DOCX (mở bằng Word, F9 để cập nhật mục lục)

Đầu ra nằm trong outputs/report/ (không commit). Thứ tự ghép là thứ tự tên file
trong docs/bao_cao/ (00_..., 01_..., ...); README.md không được ghép.
"""
from __future__ import annotations

import argparse
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "bao_cao"
OUTPUT = ROOT / "outputs" / "report"
NAME = "BaoCao_Nhom2_NASA24"
FONT = "Times New Roman"


PENDING = re.compile(r"\[CHỜ[^\]]*\]")


def chapters(only: list[str] | None = None) -> list[Path]:
    files = sorted(p for p in SOURCE.glob("[0-9][0-9]_*.md"))
    if only:
        files = [p for p in files if p.name[:2] == "00" or p.name[:2] in only]
    if not files:
        raise SystemExit(f"Không có chương nào trong {SOURCE.relative_to(ROOT)}.")
    return files


def report_status(files: list[Path]) -> None:
    """In số chữ và các chỗ còn chờ kết quả ([CHỜ ...]) của từng file để nhóm theo dõi."""
    total = 0
    for path in files:
        text = re.sub(r"<!--.*?-->", "", path.read_text(encoding="utf-8"), flags=re.S)
        words = len(re.findall(r"\w+", text))
        total += words
        pending = PENDING.findall(text)
        print(f"  {path.name:<40} {words:>6} chữ" + (f" · {len(pending)} chỗ [CHỜ]" if pending else ""))
    print(f"  {'Tổng':<40} {total:>6} chữ")


def _font(style, size: float | None = None, *, bold: bool | None = None, italic: bool | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(key), FONT)
    for key in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(key), None)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)


def _page_number_footer(section) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    paragraph = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            element = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): kind})
        elif text.strip() == "PAGE":
            element = run._r.makeelement(qn("w:instrText"), {qn("xml:space"): "preserve"})
            element.text = text
        else:
            element = run._r.makeelement(qn("w:t"), {})
            element.text = text
        run._r.append(element)


def reference_docx(pandoc: str, target: Path) -> Path:
    """Mẫu kiểu chữ cho pandoc: lấy mẫu mặc định rồi chỉnh theo quy định trình bày."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: F401
    from docx.shared import Cm, Pt

    raw = subprocess.run([pandoc, "--print-default-data-file", "reference.docx"], check=True, capture_output=True).stdout
    target.write_bytes(raw)
    document = Document(str(target))
    styles = document.styles
    body = ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text", "Footnote Text")
    for name in body:
        if name in styles:
            _font(styles[name], 13)
            fmt = styles[name].paragraph_format
            fmt.line_spacing = 1.3
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(6)
    for name in ("Body Text", "First Paragraph"):
        styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        styles[name].paragraph_format.first_line_indent = Cm(1)
    styles["Compact"].paragraph_format.space_after = Pt(2)
    styles["Block Text"].paragraph_format.left_indent = Cm(1)
    for level, size in ((1, 16), (2, 14), (3, 13), (4, 13)):
        style = styles[f"Heading {level}"]
        _font(style, size, bold=True, italic=(level >= 3))
        fmt = style.paragraph_format
        fmt.line_spacing = 1.3
        fmt.space_before = Pt(18 if level == 1 else 12)
        fmt.space_after = Pt(12 if level == 1 else 6)
        fmt.keep_with_next = True
    styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Heading 1"].paragraph_format.page_break_before = True
    for name, size in (("Title", 18), ("Subtitle", 15), ("Author", 13), ("Date", 13), ("Abstract", 13)):
        if name in styles:
            _font(styles[name], size, bold=(name == "Title"))
            styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name in ("Image Caption", "Table Caption", "Caption"):
        if name in styles:
            _font(styles[name], 12, italic=True)
            styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            styles[name].paragraph_format.space_after = Pt(10)
    for name in ("Captioned Figure", "Figure"):
        if name in styles:
            styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            styles[name].paragraph_format.keep_with_next = True
    if "Source Code" in styles:
        styles["Source Code"].font.size = Pt(10)
    # Kiểu cho trang bìa, dùng trong Markdown bằng ::: {custom-style="Cover Title"} ... :::
    from docx.enum.style import WD_STYLE_TYPE
    for name, size, bold, space in (("Cover", 13, False, 4), ("Cover Strong", 14, True, 4),
                                    ("Cover Title", 18, True, 12), ("Cover Big", 22, True, 18)):
        style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles["Normal"]
        _font(style, size, bold=bold)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_after = Pt(space)
    if "TOC Heading" in styles:
        _font(styles["TOC Heading"], 16, bold=True)
        styles["TOC Heading"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for section in document.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.left_margin, section.right_margin = Cm(3), Cm(2)
        section.top_margin, section.bottom_margin = Cm(2), Cm(2)
        section.footer_distance = Cm(1)
        _page_number_footer(section)
    document.save(str(target))
    return target


def build_docx(pandoc: str, output: Path, files: list[Path]) -> Path:
    reference = reference_docx(pandoc, output.parent / "reference.docx")
    command = [
        pandoc, *[str(p) for p in files],
        "--from", "markdown+pipe_tables+grid_tables+tex_math_dollars+implicit_figures+fenced_divs+raw_attribute",
        "--to", "docx",
        "--reference-doc", str(reference),
        "--resource-path", os.pathsep.join(str(p) for p in (SOURCE, ROOT / "docs", ROOT)),
        "--syntax-highlighting=none",  # mã in đen trắng, hợp với bản in
        "--citeproc", "--bibliography", str(SOURCE / "tai_lieu.bib"), "--csl", str(SOURCE / "ieee.csl"),
        # Danh mục tài liệu theo IEEE bằng tiếng Anh (vol., no., pp.); ngôn ngữ văn bản đặt lại vi-VN bên dưới.
        "--metadata", "lang=en-US", "--metadata", "link-citations=true",
        "--output", str(output),
    ]
    subprocess.run(command, check=True, cwd=SOURCE)
    _postprocess(output, "vi-VN")
    return output


def _postprocess(path: Path, language: str) -> None:
    """Sau pandoc: đặt ngôn ngữ văn bản (kiểm chính tả trong Word) và trình bày lại các bảng."""
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(str(path))
    defaults = document.styles.element.find(qn("w:docDefaults"))
    rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr")) if defaults is not None else None
    if rpr is not None:
        lang = rpr.find(qn("w:lang"))
        if lang is None:
            lang = rpr.makeelement(qn("w:lang"), {})
            rpr.append(lang)
        lang.set(qn("w:val"), language)
    for table in document.tables:
        _polish_table(table)
        following = table._tbl.getnext()
        if following is not None and following.tag == qn("w:p"):  # chừa khoảng giữa bảng và đoạn sau
            spacing = following.get_or_add_pPr().find(qn("w:spacing"))
            if spacing is None:
                spacing = following.get_or_add_pPr().makeelement(qn("w:spacing"), {})
                following.get_or_add_pPr().append(spacing)
            spacing.set(qn("w:before"), "160")
    document.save(str(path))


def _polish_table(table, text_width_cm: float = 16.0) -> None:
    """Bảng rộng hết dòng, cột chia theo lượng chữ, chữ 12 căn trái, hàng đầu lặp lại khi sang trang.

    Pandoc chia cột đều và để ô dùng kiểu đoạn văn (căn đều, thụt dòng), nên bảng nhiều chữ rất khó đọc.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    rows = table.rows
    columns = len(table.columns)
    if not rows or not columns:
        return
    texts = [[cell.text.strip() for cell in row.cells[:columns]] for row in rows]
    weight, longest = [], []
    for i in range(columns):
        column = [row[i] for row in texts if i < len(row)]
        weight.append(min(60.0, max(6.0, sum(len(t) for t in column) / len(column))))
        longest.append(max((len(word) for t in column for word in t.split()), default=1))
    widths = [text_width_cm * w / sum(weight) for w in weight]
    floor = [min(5.0, 0.22 * n + 0.4) for n in longest]  # khoảng 0,22 cm mỗi ký tự ở cỡ 12
    for _ in range(3):
        short = [max(f - w, 0.0) for w, f in zip(widths, floor)]
        if not any(short):
            break
        spare = [i for i in range(columns) if not short[i]]
        take = sum(short)
        room = sum(widths[i] - floor[i] for i in spare) or 1.0
        widths = [max(w, f) if short[i] else w - take * (w - floor[i]) / room for i, (w, f) in enumerate(zip(widths, floor))]
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    for tag, attrs in (("w:tblW", {"w:type": "dxa", "w:w": str(int(text_width_cm * 567))}),
                       ("w:tblLayout", {"w:type": "fixed"})):
        element = tbl_pr.find(qn(tag))
        if element is None:
            element = tbl_pr.makeelement(qn(tag), {})
            tbl_pr.append(element)
        for key, value in attrs.items():
            element.set(qn(key), value)
    grid = tbl.tblGrid
    for col, width in zip(grid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(width * 567)))
    for r, row in enumerate(rows):
        if r == 0:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:tblHeader")) is None:
                tr_pr.append(tr_pr.makeelement(qn("w:tblHeader"), {}))
        for i, cell in enumerate(row.cells[:columns]):
            cell.width = Cm(widths[i])
            for paragraph in cell.paragraphs:
                fmt = paragraph.paragraph_format
                fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
                fmt.first_line_indent = Cm(0)
                fmt.line_spacing = 1.15
                fmt.space_before, fmt.space_after = Pt(2), Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(12)
                    if r == 0:
                        run.font.bold = True


def uno_python() -> str | None:
    """Python có module uno (LibreOffice) để cập nhật mục lục rồi xuất PDF."""
    candidates = [os.environ.get("REPORT_UNO_PYTHON"), "/usr/bin/python3",
                  r"C:\Program Files\LibreOffice\program\python.exe",
                  "/Applications/LibreOffice.app/Contents/Resources/python"]
    for candidate in filter(None, candidates):
        if Path(candidate).exists() and subprocess.run([candidate, "-c", "import uno"],
                                                       capture_output=True).returncode == 0:
            return candidate
    return None


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--no-pdf", action="store_true", help="chỉ xuất DOCX")
    parser.add_argument("--only", nargs="+", metavar="NN", help="chỉ ghép trang đầu và các chương có số này, ví dụ --only 02 03")
    parser.add_argument("--out", type=Path, default=OUTPUT, help="thư mục đầu ra (mặc định outputs/report)")
    args = parser.parse_args(argv)
    files = chapters(args.only)
    report_status(files)
    try:
        import pypandoc
    except ImportError:
        raise SystemExit("Thiếu pypandoc: pip install -r requirements-report.txt")
    pandoc = pypandoc.get_pandoc_path()
    args.out.mkdir(parents=True, exist_ok=True)
    docx = build_docx(pandoc, args.out / f"{NAME}.docx", files)
    print(f"DOCX: {docx}")
    if args.no_pdf:
        return 0
    python = uno_python()
    if python is None or shutil.which("soffice") is None and os.name != "nt":
        print("Không tìm thấy LibreOffice (soffice + uno); mở DOCX bằng Word, F9 cập nhật mục lục rồi lưu PDF.")
        return 0
    pdf = args.out / f"{NAME}.pdf"
    subprocess.run([python, str(ROOT / "scripts" / "report_to_pdf.py"), str(docx), str(pdf)], check=True)
    print(f"PDF: {pdf}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
